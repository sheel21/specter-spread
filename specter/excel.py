"""Banker-formatted Excel workbook (openpyxl): Summary, one sheet per statement, Checks, Sources.

Conventions: blue font = hardcoded input from a filing; black = formula; yellow fill = flagged by
the Checker (with the note as a cell comment). Derived lines, ratios and check rows are real Excel
formulas. Content is always neutral (no persona).
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from specter import template

NAVY, GOLD, BLUE = "0B1F3A", "C5A253", "0000FF"
F_TITLE = Font(name="Georgia", size=15, bold=True, color=NAVY)
F_META = Font(name="Calibri", size=9, italic=True, color="555555")
F_HEAD = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name="Calibri", size=10, color=BLUE)
F_INPUT_B = Font(name="Calibri", size=10, color=BLUE, bold=True)
F_FORMULA = Font(name="Calibri", size=10, color="000000")
F_FORMULA_B = Font(name="Calibri", size=10, color="000000", bold=True)
F_LABEL = Font(name="Calibri", size=10)
F_LABEL_B = Font(name="Calibri", size=10, bold=True)
F_NA = Font(name="Calibri", size=10, color="999999")
F_SECTION = Font(name="Georgia", size=11, bold=True, color=NAVY)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_FLAG = PatternFill("solid", fgColor="FFF3B0")
FILL_SECTION = PatternFill("solid", fgColor="F4F1E8")
TOP = Border(top=Side(style="thin", color="000000"))
GOLD_BOTTOM = Border(bottom=Side(style="medium", color=GOLD))
RIGHT = Alignment(horizontal="right")

FMT = {
    "usd": '#,##0_);(#,##0);"–"_)',
    "shares": '#,##0_);(#,##0);"–"_)',
    "per_share": '0.00_);(0.00);"–"_)',
    "pct": '0.0%_);(0.0%);"–"_)',
}
SHEETS = {
    "annual": {"income": "Income Statement", "balance": "Balance Sheet", "cashflow": "Cash Flow"},
    "quarterly": {"income": "IS (Quarterly)", "balance": "BS (Quarterly)", "cashflow": "CF (Quarterly)"},
}
# Derived lines -> Excel formula on the same column. Keys are template line keys.
FORMULAS = {
    "gross_profit": lambda r, c: f"={c}{r['revenue']}-{c}{r['cogs']}",
    "ebitda": lambda r, c: f"={c}{r['ebit']}+{c}{r['da']}",
    "fcf": lambda r, c: f"={c}{r['cfo']}+{c}{r['capex']}",
    "total_liabilities": lambda r, c: f"={c}{r['total_liabilities_and_equity']}-{c}{r['total_equity']}",
}
FIRST_ROW = 7


def _scaled(v, unit):
    return v if unit in ("per_share", "pct") else v / 1e6


def _header(ws, spread: dict, title: str, ncols: int) -> None:
    c = spread["company"]
    ws["A1"] = f"{c['name']}: {title}"
    ws["A1"].font = F_TITLE
    ws["A2"] = (f"Ticker {c['ticker']} · CIK {c['cik']} · {spread['units']} · Source: SEC EDGAR XBRL "
                f"(10-K / 10-Q), {spread['data_source']} · Generated {spread['generated_at']}")
    ws["A3"] = "Blue = hardcoded from filings · Black = formula · Yellow = flagged by the Checker (see comment)"
    ws["A2"].font = ws["A3"].font = F_META
    for col in range(1, ncols + 1):
        ws.cell(row=4, column=col).border = GOLD_BOTTOM


def _statement_sheet(wb, spread: dict, view_name: str, statement: str) -> dict:
    view = spread["views"][view_name]
    periods = view["periods"]
    ws = wb.create_sheet(SHEETS[view_name][statement])
    _header(ws, spread, f"{template.STATEMENTS[statement]} ({view_name})", len(periods) + 1)
    ws.cell(row=5, column=1, value="($ in millions, except per share)").font = F_HEAD
    for j, p in enumerate(periods, 2):
        ws.cell(row=5, column=j, value=p["label"]).font = F_HEAD
        ws.cell(row=6, column=j, value=f"ended {p['end']}").font = F_META
        ws.cell(row=5, column=j).alignment = ws.cell(row=6, column=j).alignment = RIGHT
    for j in range(1, len(periods) + 2):
        ws.cell(row=5, column=j).fill = FILL_HEAD

    rows = view["statements"][statement]
    row_of = {r["key"]: FIRST_ROW + i for i, r in enumerate(rows)}
    for i, r in enumerate(rows):
        n = FIRST_ROW + i
        bold = r["kind"] == "subtotal"
        lab = ws.cell(row=n, column=1, value=r["label"])
        lab.font = F_LABEL_B if bold else Font(name="Calibri", size=10, italic=r["kind"] == "memo")
        for j, p in enumerate(periods, 2):
            col = get_column_letter(j)
            cell = ws.cell(row=n, column=j)
            c = r["cells"].get(p["key"]) or {}
            v, src = c.get("v"), c.get("src") or {}
            if v is None:
                cell.value, cell.font, cell.alignment = "N/A", F_NA, RIGHT
            elif src.get("type") == "derived" and r["key"] in FORMULAS:
                cell.value = FORMULAS[r["key"]](row_of, col)
                cell.font = F_FORMULA_B if bold else F_FORMULA
            else:
                cell.value = _scaled(v, r["unit"])
                cell.font = F_INPUT_B if bold else F_INPUT
            cell.number_format = FMT[r["unit"]]
            if bold:
                cell.border = TOP
            if c.get("flag"):
                cell.fill = FILL_FLAG
                cell.comment = Comment(c["flag"][:1500], "Checker")
        if bold:
            lab.border = TOP

    n = FIRST_ROW + len(rows) + 1
    if statement == "income":
        n = _ratio_block(ws, view, periods, row_of, n)
    n = _check_rows(ws, statement, periods, row_of, n)

    ws.column_dimensions["A"].width = 40
    for j in range(2, len(periods) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.freeze_panes = "B7"
    ws.sheet_view.showGridLines = False
    return row_of


def _ratio_block(ws, view, periods, row_of, n) -> int:
    ws.cell(row=n, column=1, value="Growth & Margins").font = F_SECTION
    for j in range(1, len(periods) + 2):
        ws.cell(row=n, column=j).fill = FILL_SECTION
    n += 1
    lag = 4 if periods and periods[0]["fq"] else 1  # quarters grow year over year
    ratios = {r["key"]: r for r in view["ratios"]}
    spec = [("revenue_growth", None), ("gross_margin", "gross_profit"), ("ebit_margin", "ebit"),
            ("ebitda_margin", "ebitda"), ("net_margin", "net_income")]
    for key, numerator in spec:
        r = ratios[key]
        ws.cell(row=n, column=1, value=r["label"]).font = Font(name="Calibri", size=10, italic=True)
        for j, p in enumerate(periods, 2):
            col, cell = get_column_letter(j), ws.cell(row=n, column=j)
            rev = row_of["revenue"]
            if numerator:
                cell.value = f'=IFERROR({col}{row_of[numerator]}/{col}{rev},"N/A")'
                cell.font = F_FORMULA
            elif j - lag >= 2:
                prev = get_column_letter(j - lag)
                cell.value = f'=IFERROR({col}{rev}/{prev}{rev}-1,"N/A")'
                cell.font = F_FORMULA
            else:  # prior period isn't on the sheet: computed from the filed prior-period revenue
                v = (r["cells"].get(p["key"]) or {}).get("v")
                cell.value, cell.font = (v, F_INPUT) if v is not None else ("N/A", F_NA)
            cell.number_format = FMT["pct"]
            cell.alignment = RIGHT
        n += 1
    return n + 1


def _check_rows(ws, statement, periods, row_of, n) -> int:
    """Live check rows: they should read 0 (shown as a dash)."""
    checks = {
        "balance": [("Check: Assets − (Liabilities + Equity)",
                     lambda c: f"={c}{row_of['total_assets']}-({c}{row_of['total_liabilities']}+{c}{row_of['total_equity']})")],
        "income": [("Check: Gross Profit − (Revenue − COGS)",
                    lambda c: f"={c}{row_of['gross_profit']}-({c}{row_of['revenue']}-{c}{row_of['cogs']})")],
        "cashflow": [("Check: CFO + CFI + CFF + FX − Net Change",
                      lambda c: f"={c}{row_of['cfo']}+{c}{row_of['cfi']}+{c}{row_of['cff']}+N({c}{row_of['fx_effect']})"
                                f"-{c}{row_of['net_change_cash']}")],
    }[statement]
    for label, formula in checks:
        ws.cell(row=n, column=1, value=label).font = F_META
        for j in range(2, len(periods) + 2):
            col, cell = get_column_letter(j), ws.cell(row=n, column=j)
            cell.value = f'=IFERROR({formula(col)[1:]},"N/A")'
            cell.font, cell.number_format = F_META, FMT["usd"]
        n += 1
    return n


def _summary(wb, spread, refs) -> None:
    ws = wb.active
    ws.title = "Summary"
    c = spread["company"]
    _header(ws, spread, "Summary", 8)
    view_name = "annual" if "annual" in spread["views"] else "quarterly"
    view, sheet = spread["views"][view_name], SHEETS[view_name]
    periods = view["periods"]
    ws.cell(row=5, column=1, value="Key metrics ($ in millions)").font = F_HEAD
    for j, p in enumerate(periods, 2):
        ws.cell(row=5, column=j, value=p["label"]).font = F_HEAD
        ws.cell(row=5, column=j).alignment = RIGHT
    for j in range(1, len(periods) + 2):
        ws.cell(row=5, column=j).fill = FILL_HEAD
    ratio_row = {k: FIRST_ROW + len(view["statements"]["income"]) + 2 + i
                 for i, k in enumerate(["revenue_growth", "gross_margin", "ebit_margin", "ebitda_margin", "net_margin"])}
    lines = [("Revenue", sheet["income"], refs[view_name]["income"]["revenue"], "usd"),
             ("Revenue Growth %", sheet["income"], ratio_row["revenue_growth"], "pct"),
             ("Gross Margin %", sheet["income"], ratio_row["gross_margin"], "pct"),
             ("EBITDA", sheet["income"], refs[view_name]["income"]["ebitda"], "usd"),
             ("EBITDA Margin %", sheet["income"], ratio_row["ebitda_margin"], "pct"),
             ("Net Income", sheet["income"], refs[view_name]["income"]["net_income"], "usd"),
             ("Free Cash Flow", sheet["cashflow"], refs[view_name]["cashflow"]["fcf"], "usd"),
             ("Total Assets", sheet["balance"], refs[view_name]["balance"]["total_assets"], "usd")]
    for i, (label, sh, row, unit) in enumerate(lines):
        n = 6 + i
        ws.cell(row=n, column=1, value=label).font = F_LABEL_B if unit == "usd" else F_LABEL
        for j in range(2, len(periods) + 2):
            cell = ws.cell(row=n, column=j, value=f"='{sh}'!{get_column_letter(j)}{row}")
            cell.font, cell.number_format = F_FORMULA, FMT[unit]
    n = 6 + len(lines) + 1
    ws.cell(row=n, column=1, value="Checker").font = F_SECTION
    for name, res in spread.get("checks", {}).items():
        n += 1
        k = res["counts"]
        ws.cell(row=n, column=1, value=f"{name.title()}: {k['passed']} passed · {k['fixed']} fixed · {k['flagged']} flagged")
    n += 2
    ws.cell(row=n, column=1, value="Filings used").font = F_SECTION
    for f in spread["filings"][:20]:
        n += 1
        ws.cell(row=n, column=1, value=f"{f['form']} for period ended {f['report_date']} (filed {f['filing_date']})")
        link = ws.cell(row=n, column=2, value=f["accession"])
        link.hyperlink, link.font = f["url"], Font(color=BLUE, underline="single", size=10)
    ws.column_dimensions["A"].width = 52
    for j in range(2, len(periods) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.freeze_panes = "B6"
    ws.sheet_view.showGridLines = False
    _ = c


def _checks_sheet(wb, spread) -> None:
    ws = wb.create_sheet("Checks")
    _header(ws, spread, "Checker results", 6)
    heads = ["View", "Period", "Check", "Status", "Detail", "Fix attempts / note"]
    for j, h in enumerate(heads, 1):
        cell = ws.cell(row=5, column=j, value=h)
        cell.font, cell.fill = F_HEAD, FILL_HEAD
    n = 6
    order = {"flagged": 0, "fixed": 1, "passed": 2}
    for view_name, res in spread.get("checks", {}).items():
        for r in sorted(res["results"], key=lambda r: order[r["status"]]):
            attempts = "; ".join(f"#{a['attempt']}: {a.get('line', 'none')} via {a.get('hint', '-')} "
                                 f"({'passed' if a.get('passed') else 'failed'})" for a in r["attempts"])
            vals = [view_name, r["period_label"], r["check"], r["status"].upper(), r["detail"],
                    " ".join(x for x in (attempts, r["note"]) if x)]
            for j, v in enumerate(vals, 1):
                cell = ws.cell(row=n, column=j, value=v)
                cell.font = F_LABEL
                if r["status"] == "flagged":
                    cell.fill = FILL_FLAG
            n += 1
    for col, w in zip("ABCDEF", (10, 10, 44, 10, 80, 80)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A6"


def _filed_parts(src: dict) -> list[dict]:
    """Every filed fact behind a cell (a sum's parts can themselves be YTD differences)."""
    if src.get("type") == "xbrl":
        return [src]
    return [x for part in src.get("parts") or [] for x in _filed_parts(part)]


def _sources_sheet(wb, spread) -> None:
    ws = wb.create_sheet("Sources")
    _header(ws, spread, "Sources: XBRL tag and filing for every number", 9)
    heads = ["View", "Statement", "Line", "Period", "Value ($M)", "How obtained", "XBRL tag(s)", "Form · Accession · Filed",
             "Filing link"]
    for j, h in enumerate(heads, 1):
        cell = ws.cell(row=5, column=j, value=h)
        cell.font, cell.fill = F_HEAD, FILL_HEAD
    n = 6
    how = {"xbrl": "Filed value", "sum": "Sum of filed components", "ytd_difference": "Quarter = YTD difference",
           "derived": "Formula on this spread"}
    for view_name, view in spread["views"].items():
        for st, rows in view["statements"].items():
            for r in rows:
                for p in view["periods"]:
                    c = r["cells"].get(p["key"]) or {}
                    v, src = c.get("v"), c.get("src") or {}
                    parts = _filed_parts(src)
                    filing = "; ".join(f"{x['form']} · {x['accn']} · {x['filed']}" for x in parts)
                    vals = [view_name, template.STATEMENTS[st], r["label"], p["label"],
                            None if v is None else _scaled(v, r["unit"]),
                            "Not filed (N/A)" if v is None else how.get(src.get("type"), src.get("type")) +
                            (" · re-extracted by Checker" if src.get("re_extracted_by") else ""),
                            src.get("tag") or src.get("formula") or "", filing, ""]
                    for j, val in enumerate(vals, 1):
                        ws.cell(row=n, column=j, value=val).font = F_LABEL
                    ws.cell(row=n, column=5).number_format = FMT["per_share" if r["unit"] == "per_share" else "usd"]
                    if parts:
                        link = ws.cell(row=n, column=9, value="open")
                        link.hyperlink, link.font = parts[0]["url"], Font(color=BLUE, underline="single", size=10)
                    if c.get("flag"):
                        for j in range(1, 10):
                            ws.cell(row=n, column=j).fill = FILL_FLAG
                    n += 1
    for col, w in zip("ABCDEFGHI", (10, 18, 34, 10, 12, 30, 70, 46, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:I{n - 1}"


def build_workbook(spread: dict) -> bytes:
    wb = Workbook()
    refs = {}
    for view_name in ("annual", "quarterly"):
        if view_name in spread["views"]:
            refs[view_name] = {st: _statement_sheet(wb, spread, view_name, st) for st in ("income", "balance", "cashflow")}
    _summary(wb, spread, refs)
    _checks_sheet(wb, spread)
    _sources_sheet(wb, spread)
    wb.move_sheet("Summary", offset=-wb.index(wb["Summary"]))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
